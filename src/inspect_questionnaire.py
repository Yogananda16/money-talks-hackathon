import openpyxl

wb = openpyxl.load_workbook("questions/Regodit_Comprehensive_Vendor_Security_Questionnaire_Clean.xlsx", data_only=True)

for sheet in wb.worksheets:
    print(f"\n=== Sheet: {sheet.title} ({sheet.max_row} rows, {sheet.max_column} cols) ===")
    headers = [cell.value for cell in sheet[1]]
    print("Headers:", headers)
    print("Row 2:", [cell.value for cell in sheet[2]])
    print("Row 3:", [cell.value for cell in sheet[3]])